import json
import os
import secrets
import tempfile
from io import BytesIO
from urllib.parse import quote
from datetime import timedelta

from django.conf import settings
try:
    from asgiref.sync import async_to_sync
    from channels.layers import get_channel_layer
except ImportError:  # pragma: no cover - websocket support optional locally
    async_to_sync = None
    get_channel_layer = None
from django.contrib import messages
from django.contrib.auth import login
from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.contrib.auth.mixins import LoginRequiredMixin, UserPassesTestMixin
from django.core.mail import send_mail
from django.db.models import Max, Q, Sum
from django.db import transaction
from django.db.models import Count, Prefetch
from django.http import FileResponse, Http404, HttpResponse, HttpResponseForbidden, JsonResponse
from django.shortcuts import get_object_or_404, redirect
from django.template.loader import render_to_string
from django.urls import reverse, reverse_lazy
from django.utils import translation
from django.utils import timezone
#from django.utils.translation import LANGUAGE_SESSION_KEY
from django.views import View
from django.views.decorators.csrf import csrf_exempt
from django.views.decorators.http import require_POST
from django.views.generic import CreateView, DetailView, ListView, TemplateView, UpdateView
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont

from .forms import AdminSupportReplyForm, ContactAdminForm, ContactAdminReplyForm, ExcelUploadForm, InvitationForm, PaymentInitiationForm, SignUpForm, SiteRatingForm, WelcomeMessageForm
from .localization import tr_text
from .models import Invitation, OrganizerProfile, PaymentTransaction, SiteRating, SiteVisit, Subscription, SupportMessage

LANGUAGE_SESSION_KEY = "django_language"


def _is_ajax_request(request):
    return request.headers.get("X-Requested-With") == "XMLHttpRequest"


def _serialize_support_message(message, for_staff=False):
    delivery_state = "read" if (
        (message.sender_type == SupportMessage.SENDER_ADMIN and message.is_read_by_user)
        or (message.sender_type == SupportMessage.SENDER_USER and message.is_read_by_admin)
    ) else "delivered"
    return {
        "id": message.id,
        "organizer_id": message.organizer_id,
        "organizer_username": message.organizer.user.username,
        "organizer_display_name": message.organizer.user.get_full_name() or message.organizer.user.username,
        "sender_type": message.sender_type,
        "sender_label": (
            "Administration"
            if message.sender_type == SupportMessage.SENDER_ADMIN
            else (message.organizer.user.username if for_staff else tr_text("Vous", "You"))
        ),
        "subject": message.subject,
        "message": message.message,
        "created_at": timezone.localtime(message.created_at).strftime("%d/%m/%Y %H:%M"),
        "created_at_iso": message.created_at.isoformat(),
        "delivery_state": delivery_state,
        "delivery_label": tr_text("Lu", "Read") if delivery_state == "read" else tr_text("Distribue", "Delivered"),
        "is_read_by_user": message.is_read_by_user,
        "is_read_by_admin": message.is_read_by_admin,
    }


def _broadcast_group(group_name, payload):
    if async_to_sync is None or get_channel_layer is None:
        return
    layer = get_channel_layer()
    if layer is None:
        return
    async_to_sync(layer.group_send)(
        group_name,
        {
            "type": "support.event",
            "payload": payload,
        },
    )


def _broadcast_support_message(message):
    unread_for_user = message.organizer.support_messages.filter(
        sender_type=SupportMessage.SENDER_ADMIN,
        is_read_by_user=False,
    ).count()
    unread_for_admin_thread = message.organizer.support_messages.filter(
        sender_type=SupportMessage.SENDER_USER,
        is_read_by_admin=False,
    ).count()
    unread_for_admin_total = SupportMessage.objects.filter(
        sender_type=SupportMessage.SENDER_USER,
        is_read_by_admin=False,
    ).count()
    latest_timestamp = message.created_at.isoformat()

    _broadcast_group(
        f"support_user_{message.organizer_id}",
        {
            "kind": "support.message",
            "unread_count": unread_for_user,
            "latest_timestamp": latest_timestamp,
            "message_item": _serialize_support_message(message, for_staff=False),
        },
    )
    _broadcast_group(
        "support_staff",
        {
            "kind": "support.message",
            "unread_count": unread_for_admin_total,
            "thread_unread_count": unread_for_admin_thread,
            "latest_timestamp": latest_timestamp,
            "message_item": _serialize_support_message(message, for_staff=True),
        },
    )


def _broadcast_support_permissions(profile):
    _broadcast_group(
        f"support_user_{profile.id}",
        {
            "kind": "support.permissions",
            "can_send_message": profile.support_user_can_send,
        },
    )


def _broadcast_support_read_state(profile, read_by_admin_ids=None, read_by_user_ids=None):
    staff_payload = {
        "kind": "support.read",
        "organizer_id": profile.id,
        "thread_unread_count": profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_USER,
            is_read_by_admin=False,
        ).count(),
        "user_unread_count": profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_ADMIN,
            is_read_by_user=False,
        ).count(),
        "unread_count": SupportMessage.objects.filter(
            sender_type=SupportMessage.SENDER_USER,
            is_read_by_admin=False,
        ).count(),
        "read_by_admin_ids": read_by_admin_ids or [],
        "read_by_user_ids": read_by_user_ids or [],
    }
    user_payload = {
        "kind": "support.read",
        "organizer_id": profile.id,
        "thread_unread_count": profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_USER,
            is_read_by_admin=False,
        ).count(),
        "unread_count": profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_ADMIN,
            is_read_by_user=False,
        ).count(),
        "read_by_admin_ids": read_by_admin_ids or [],
        "read_by_user_ids": read_by_user_ids or [],
    }
    _broadcast_group(
        "support_staff",
        staff_payload,
    )
    _broadcast_group(
        f"support_user_{profile.id}",
        user_payload,
    )


def _mark_admin_messages_read_by_user(profile):
    unread_ids = list(
        profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_ADMIN,
            is_read_by_user=False,
        ).values_list("id", flat=True)
    )
    if unread_ids:
        profile.support_messages.filter(id__in=unread_ids).update(is_read_by_user=True)
        _broadcast_support_read_state(profile, read_by_user_ids=unread_ids)
    return unread_ids



def _combine_limits(previous_subscription, new_limit):
    if not previous_subscription or previous_subscription.invitation_limit is None or new_limit is None:
        return None
    return previous_subscription.remaining_invitations + new_limit


def _build_subscription(profile, provider, invite_count=None, previous_subscription=None):
    invite_count = invite_count or profile.planned_invitations
    plan_code, invitation_limit, price_usd = Subscription.pick_plan(invite_count)
    combined_limit = _combine_limits(previous_subscription, invitation_limit)
    return Subscription.objects.create(
        organizer=profile,
        plan_code=plan_code,
        invitation_limit=combined_limit if previous_subscription else invitation_limit,
        price_usd=price_usd,
        provider=provider,
        status=Subscription.STATUS_PENDING,
    )


def _mobile_money_gateway_payload(transaction):
    return {
        "merchant_reference": transaction.transaction_reference,
        "amount_usd": str(transaction.amount_usd),
        "currency": "USD",
        "provider": transaction.provider,
        "phone_number": transaction.phone_number,
        "customer_username": transaction.subscription.organizer.user.username,
    }


def _activate_subscription(subscription):
    previous_active = (
        subscription.organizer.subscriptions.filter(status=Subscription.STATUS_ACTIVE)
        .exclude(pk=subscription.pk)
        .first()
    )
    subscription.is_paid = True
    subscription.status = Subscription.STATUS_ACTIVE
    subscription.paid_at = timezone.now()
    subscription.save(update_fields=["is_paid", "status", "paid_at"])
    if previous_active:
        previous_active.status = Subscription.STATUS_EXPIRED
        previous_active.save(update_fields=["status"])


def _provider_configuration(provider_code):
    key = provider_code.upper()
    return {
        "endpoint": os.environ.get(f"{key}_API_URL", ""),
        "api_key": os.environ.get(f"{key}_API_KEY", ""),
        "api_secret": os.environ.get(f"{key}_API_SECRET", ""),
    }


def hmac_compare(left, right):
    if not left or not right:
        return False
    return secrets.compare_digest(left, right)


def _load_font(size, bold=False):
    candidates = [
        "C:/Windows/Fonts/georgiab.ttf" if bold else "C:/Windows/Fonts/georgia.ttf",
        "C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf",
    ]
    for path in candidates:
        if os.path.exists(path):
            return ImageFont.truetype(path, size=size)
    return ImageFont.load_default()


def _default_welcome_message(invitation):
    return invitation.welcome_message.strip() or tr_text(
        "Nous sommes honorés de vous accueillir parmi nos invites.",
        "We are honored to welcome you among our guests.",
    )


def _ensure_invitation_qr(invitation):
    invitation.ensure_qr_code()
    return invitation


def _build_default_invitation_message(profile, guest_name):
    if profile.default_welcome_message.strip():
        return profile.default_welcome_message.strip()
    event_label = profile.get_event_label()
    return tr_text(
        f"{guest_name}, bienvenue a notre {event_label.lower()}. Nous serons heureux de partager ce moment avec vous.",
        f"{guest_name}, welcome to our {event_label.lower()}. We will be delighted to share this moment with you.",
    )


def build_invitation_jpeg(invitation):
    invitation = _ensure_invitation_qr(invitation)
    canvas = Image.new("RGB", (1600, 920), "#efe6d7")
    draw = ImageDraw.Draw(canvas)
    draw.rounded_rectangle((40, 40, 1560, 880), radius=34, fill="#fffaf4", outline="#d8c8b0", width=3)
    draw.rounded_rectangle((60, 60, 1210, 860), radius=28, fill="#fffdf9", outline="#d6c4a8", width=2)
    draw.rounded_rectangle((1225, 60, 1540, 860), radius=28, fill="#f6efe2", outline="#d6c4a8", width=2)
    draw.rectangle((80, 88, 1190, 150), fill="#143d64")
    draw.rectangle((1248, 88, 1518, 150), fill="#c67d2d")

    title_font = _load_font(70, bold=True)
    subtitle_font = _load_font(30, bold=True)
    body_font = _load_font(28)
    small_font = _load_font(22)
    code_font = _load_font(26, bold=True)

    draw.text((108, 98), tr_text("CARTE D'EMBARQUEMENT INVITATION", "INVITATION BOARDING PASS"), fill="#ffffff", font=subtitle_font)
    draw.text((132, 205), invitation.guest_name, fill="#1f1a17", font=title_font)
    draw.text((132, 300), invitation.event_label.upper(), fill="#143d64", font=subtitle_font)
    draw.text((132, 372), tr_text("Bienvenue a bord de cette celebration", "Welcome aboard this celebration"), fill="#6e6258", font=body_font)
    draw.text(
        (132, 438),
        _default_welcome_message(invitation),
        fill="#6e6258",
        font=body_font,
    )
    draw.text(
        (132, 572),
        tr_text(
            f"Organisateur: {invitation.organizer.user.get_full_name() or invitation.organizer.user.username}",
            f"Organizer: {invitation.organizer.user.get_full_name() or invitation.organizer.user.username}",
        ),
        fill="#6e6258",
        font=small_font,
    )
    draw.text((132, 628), tr_text(f"Porte / Zone: {invitation.seat_location}", f"Gate / Zone: {invitation.seat_location}"), fill="#143d64", font=code_font)
    draw.text((132, 678), tr_text(f"Code billet: {invitation.slug.upper()}", f"Ticket code: {invitation.slug.upper()}"), fill="#1f1a17", font=small_font)
    draw.text((132, 728), tr_text("QR code integre pour acces, controle et placement", "QR code embedded for access, control, and seating"), fill="#0f5d5e", font=small_font)
    draw.text((1270, 205), tr_text("ACCES", "ACCESS"), fill="#c67d2d", font=subtitle_font)
    draw.text((1270, 270), invitation.event_label.upper(), fill="#1f1a17", font=subtitle_font)
    draw.text((1270, 360), tr_text("Invité", "Guest"), fill="#6e6258", font=small_font)
    draw.text((1270, 395), invitation.guest_name[:18], fill="#1f1a17", font=body_font)
    draw.text((1270, 470), tr_text("Place", "Seat"), fill="#6e6258", font=small_font)
    draw.text((1270, 505), invitation.seat_location[:18], fill="#1f1a17", font=body_font)

    if invitation.qr_code:
        with invitation.qr_code.open("rb") as qr_file:
            qr_image = Image.open(qr_file).convert("RGB")
            qr_image = qr_image.resize((210, 210))
            canvas.paste(qr_image, (1280, 585))

    output = BytesIO()
    canvas.save(output, format="JPEG", quality=92)
    output.seek(0)
    return output


def build_staff_report_workbook():
    workbook = Workbook()
    users_sheet = workbook.active
    users_sheet.title = "Utilisateurs"
    users_sheet.append(["Username", "Nom complet", "Email", "Ceremonie", "Telephone", "Invitations prevues"])
    for profile in OrganizerProfile.objects.select_related("user").all():
        users_sheet.append(
            [
                profile.user.username,
                profile.user.get_full_name(),
                profile.user.email,
                profile.get_event_label(),
                profile.phone_number,
                profile.planned_invitations,
            ]
        )

    subscriptions_sheet = workbook.create_sheet("Abonnements")
    subscriptions_sheet.append(["Utilisateur", "Plan", "Quota", "Prix USD", "Provider", "Statut", "Paye le"])
    for subscription in Subscription.objects.select_related("organizer", "organizer__user").all():
        subscriptions_sheet.append(
            [
                subscription.organizer.user.username,
                subscription.get_plan_code_display(),
                subscription.invitation_limit or "Illimite",
                float(subscription.price_usd),
                subscription.get_provider_display(),
                subscription.get_status_display(),
                subscription.paid_at.strftime("%Y-%m-%d %H:%M") if subscription.paid_at else "",
            ]
        )
    visits_sheet = workbook.create_sheet("Visites")
    visits_sheet.append(["Page", "Session", "Utilisateur", "Hits", "Derniere visite"])
    for visit in SiteVisit.objects.select_related("user").all()[:500]:
        visits_sheet.append(
            [
                visit.path,
                visit.session_key,
                visit.user.username if visit.user else "",
                visit.hits,
                visit.last_seen_at.strftime("%Y-%m-%d %H:%M"),
            ]
        )
    return workbook


def build_staff_stats():
    now = timezone.now()
    month_start = now.replace(day=1, hour=0, minute=0, second=0, microsecond=0)
    active_cutoff = now - timezone.timedelta(days=30)
    week_cutoff = now - timedelta(days=6)

    invitations_ordered = OrganizerProfile.objects.aggregate(total=Sum("planned_invitations"))["total"] or 0
    invitations_used = Invitation.objects.count()
    active_users = OrganizerProfile.objects.filter(last_seen_at__gte=active_cutoff).count()
    inactive_users = OrganizerProfile.objects.count() - active_users
    total_visits = SiteVisit.objects.aggregate(total=Sum("hits"))["total"] or 0
    unique_visitors = SiteVisit.objects.values("session_key").distinct().count()
    monthly_revenue = (
        Subscription.objects.filter(is_paid=True, paid_at__gte=month_start).aggregate(total=Sum("price_usd"))["total"]
        or 0
    )
    daily_visits = []
    for offset in range(7):
        target_day = (week_cutoff + timedelta(days=offset)).date()
        hits = SiteVisit.objects.filter(last_seen_at__date=target_day).aggregate(total=Sum("hits"))["total"] or 0
        daily_visits.append({"label": target_day.strftime("%d/%m"), "total": hits})

    return {
        "ceremonies": list(
            OrganizerProfile.objects.values("ceremony_type").annotate(total=Count("id")).order_by("ceremony_type")
        ),
        "plans": list(
            Subscription.objects.values("plan_code").annotate(total=Count("id")).order_by("plan_code")
        ),
        "usage": [
            {"label": tr_text("Commandees", "Ordered"), "total": invitations_ordered},
            {"label": tr_text("Utilisees", "Used"), "total": invitations_used},
        ],
        "activity": [
            {"label": tr_text("Actifs", "Active"), "total": active_users},
            {"label": tr_text("Non actifs", "Inactive"), "total": max(inactive_users, 0)},
        ],
        "revenue": [
            {"label": month_start.strftime("%b %Y"), "total": float(monthly_revenue)},
        ],
        "visits": daily_visits,
        "visit_totals": {
            "total_hits": total_visits,
            "unique_visitors": unique_visitors,
        },
    }


def _connected_profiles_queryset():
    return OrganizerProfile.objects.select_related("user").filter(
        last_seen_at__gte=timezone.now() - timedelta(minutes=5)
    ).order_by("-last_seen_at")


def create_invitations_from_workbook(profile, excel_file):
    with tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx") as temp_file:
        for chunk in excel_file.chunks():
            temp_file.write(chunk)
        temp_path = temp_file.name

    try:
        workbook = load_workbook(temp_path, data_only=True)
        sheet = workbook.active
        created = 0
        skipped = 0
        subscription = profile.active_subscription
        if not subscription or not subscription.is_paid:
            return 0, 0, "Votre abonnement doit etre actif avant l'import Excel."

        remaining = subscription.remaining_invitations
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else ""
            seat = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""

            if index == 1 and name.lower() in {"nom", "invite", "nom de l'invite", "nom du couple"}:
                continue
            if not name or not seat:
                skipped += 1
                continue
            if remaining is not None and created >= remaining:
                skipped += 1
                continue

            Invitation.objects.create(
                organizer=profile,
                subscription=subscription,
                guest_name=name,
                seat_location=seat,
                welcome_message=_build_default_invitation_message(profile, name),
            )
            created += 1
        return created, skipped, ""
    finally:
        if os.path.exists(temp_path):
            os.remove(temp_path)


class HomeView(TemplateView):
    template_name = "invitations/home.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        if self.request.user.is_authenticated and hasattr(self.request.user, "organizer_profile"):
            profile = self.request.user.organizer_profile
            invitations = Invitation.objects.filter(organizer=profile)
            subscription = profile.active_subscription
            user_rating = getattr(profile, "site_rating", None)
            context.update(
                {
                    "profile": profile,
                    "subscription": subscription,
                    "invitations": invitations[:12],
                    "user_rating": user_rating.stars if user_rating else 0,
                    "rating_form": SiteRatingForm(initial={"stars": user_rating.stars if user_rating else 5}),
                    "stats": {
                        "total": invitations.count(),
                        "printed": invitations.exclude(printed_at__isnull=True).count(),
                        "shared": invitations.exclude(whatsapp_shared_at__isnull=True).count(),
                        "remaining": subscription.remaining_invitations if subscription else 0,
                    },
                }
            )
        else:
            context["plans"] = [
                {"name": "Starter", "range": tr_text("1 a 199 invitations", "1 to 199 invitations"), "price": "79,9 $"},
                {"name": "Pro", "range": tr_text("200 a 599 invitations", "200 to 599 invitations"), "price": "239,9 $"},
                {"name": tr_text("Illimite", "Unlimited"), "range": tr_text("Illimite", "Unlimited"), "price": "999,9 $"},
            ]
        return context


class AboutView(TemplateView):
    template_name = "invitations/about.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["ceo_name"] = "Tresor"
        context["ceo_phone"] = "+243974073701"
        return context


class SignUpView(CreateView):
    form_class = SignUpForm
    template_name = "registration/signup.html"
    success_url = reverse_lazy("invitations:subscription")

    def form_valid(self, form):
        self.object = form.save()
        profile = OrganizerProfile.objects.create(
            user=self.object,
            ceremony_type=form.cleaned_data["ceremony_type"],
            planned_invitations=form.cleaned_data["planned_invitations"],
            phone_number=form.cleaned_data["phone_number"],
            preferred_language=OrganizerProfile.LANG_FR,
        )
        _build_subscription(profile, form.cleaned_data["provider"])
        language = OrganizerProfile.LANG_FR
        login(self.request, self.object)
        self.request.session[LANGUAGE_SESSION_KEY] = language
        translation.activate(language)
        self.request.LANGUAGE_CODE = language
        messages.success(
            self.request,
            tr_text("Votre compte a ete cree. Finalisez maintenant votre paiement.", "Your account has been created. Complete your payment now."),
        )
        return redirect("invitations:subscription")


@require_POST
def set_language_preference(request):
    language = (request.POST.get("language") or OrganizerProfile.LANG_FR)[:2]
    if language not in {OrganizerProfile.LANG_FR, OrganizerProfile.LANG_EN}:
        language = OrganizerProfile.LANG_FR

    request.session[LANGUAGE_SESSION_KEY] = language
    translation.activate(language)

    if request.user.is_authenticated:
        profile = getattr(request.user, "organizer_profile", None)
        if profile and profile.preferred_language != language:
            profile.preferred_language = language
            profile.save(update_fields=["preferred_language"])

    next_url = request.POST.get("next") or reverse("invitations:home")
    return redirect(next_url)


class SubscriptionDashboardView(LoginRequiredMixin, TemplateView):
    template_name = "invitations/subscription_dashboard.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        profile = self.request.user.organizer_profile
        subscription = profile.active_subscription or profile.subscriptions.first()
        payments = PaymentTransaction.objects.filter(subscription__organizer=profile).select_related("subscription")
        show_payment_form = not (subscription and subscription.is_paid and not self.request.GET.get("renew"))
        context.update(
            {
                "profile": profile,
                "subscription": subscription,
                "active_subscription": profile.active_subscription,
                "payments": payments,
                "show_payment_form": show_payment_form,
                "welcome_form": WelcomeMessageForm(instance=profile),
                "payment_form": PaymentInitiationForm(
                    initial={
                        "planned_invitations": profile.planned_invitations,
                        "phone_number": profile.phone_number,
                        "provider": subscription.provider if subscription else Subscription.PROVIDER_DEMO_CARD,
                    }
                ),
            }
        )
        return context


class StaffRequiredMixin(UserPassesTestMixin):
    def test_func(self):
        return self.request.user.is_staff

    def handle_no_permission(self):
        messages.error(self.request, "Cette page est reservee a l'administration.")
        return redirect("invitations:home")


class InvitationPermissionMixin(LoginRequiredMixin):
    model = Invitation

    def get_queryset(self):
        return Invitation.objects.filter(organizer=self.request.user.organizer_profile)

    def dispatch(self, request, *args, **kwargs):
        profile = getattr(request.user, "organizer_profile", None)
        if not profile:
            messages.error(request, "Votre compte organisateur est incomplet.")
            return redirect("invitations:signup")
        if not profile.has_paid_access:
            messages.error(request, "Un abonnement paye est necessaire pour acceder aux invitations.")
            return redirect("invitations:subscription")
        return super().dispatch(request, *args, **kwargs)


class InvitationListView(InvitationPermissionMixin, ListView):
    template_name = "invitations/invitation_list.html"
    context_object_name = "invitations"
    paginate_by = 12

    def get_queryset(self):
        return super().get_queryset().select_related("subscription", "organizer", "organizer__user")

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["subscription"] = self.request.user.organizer_profile.active_subscription
        context["upload_form"] = ExcelUploadForm()
        return context


class InvitationExcelImportView(InvitationPermissionMixin, View):
    @transaction.atomic
    def post(self, request):
        form = ExcelUploadForm(request.POST, request.FILES)
        if not form.is_valid():
            messages.error(request, "Le fichier Excel est invalide. Utilisez un fichier .xlsx.")
            return redirect("invitations:list")

        created, skipped, error_message = create_invitations_from_workbook(
            request.user.organizer_profile,
            form.cleaned_data["excel_file"],
        )
        if error_message:
            messages.error(request, error_message)
            return redirect("invitations:subscription")

        if created:
            messages.success(
                request,
                f"Import termine: {created} invitation(s) creee(s) automatiquement.",
            )
        if skipped:
            messages.warning(
                request,
                f"{skipped} ligne(s) ont ete ignoree(s) car incomplètes ou hors quota.",
            )
        return redirect("invitations:list")


class InvitationCreateView(InvitationPermissionMixin, CreateView):
    form_class = InvitationForm
    template_name = "invitations/invitation_form.html"

    @transaction.atomic
    def form_valid(self, form):
        profile = OrganizerProfile.objects.select_for_update().get(user=self.request.user)
        subscription = profile.active_subscription
        if not subscription or not subscription.is_paid:
            messages.error(self.request, "Votre abonnement doit etre paye avant de creer une invitation.")
            return redirect("invitations:subscription")

        if subscription.invitation_limit is not None and profile.invitations.count() >= subscription.invitation_limit:
            messages.error(self.request, "Votre quota d'invitations a ete atteint.")
            return redirect("invitations:list")

        form.instance.organizer = profile
        form.instance.subscription = subscription
        if not form.instance.welcome_message.strip():
            form.instance.welcome_message = _build_default_invitation_message(profile, form.instance.guest_name)
        messages.success(self.request, "Invitation creee avec succes.")
        return super().form_valid(form)

    def get_success_url(self):
        return reverse("invitations:detail", kwargs={"slug": self.object.slug})


class InvitationDetailView(InvitationPermissionMixin, DetailView):
    template_name = "invitations/invitation_detail.html"
    context_object_name = "invitation"

    def get_context_data(self, **kwargs):
        _ensure_invitation_qr(self.object)
        context = super().get_context_data(**kwargs)
        context["can_edit"] = not self.object.is_locked
        context["download_enabled"] = self.object.subscription.is_paid
        absolute_url = self.request.build_absolute_uri(
            reverse("invitations:shared-detail", kwargs={"token": self.object.share_token})
        )
        text = (
            f"Invitation {self.object.event_label} pour {self.object.guest_name}. "
            f"Place: {self.object.seat_location}. Voir l'invitation: {absolute_url}"
        )
        context["whatsapp_url"] = f"https://wa.me/?text={quote(text)}"
        return context


class InvitationUpdateView(InvitationPermissionMixin, UpdateView):
    form_class = InvitationForm
    template_name = "invitations/invitation_form.html"

    def dispatch(self, request, *args, **kwargs):
        invitation = self.get_object()
        if invitation.is_locked:
            messages.error(
                request,
                "Cette invitation ne peut plus etre modifiee apres impression ou partage WhatsApp.",
            )
            return redirect("invitations:detail", slug=invitation.slug)
        return super().dispatch(request, *args, **kwargs)

    def get_success_url(self):
        if not self.object.welcome_message.strip():
            self.object.welcome_message = _build_default_invitation_message(self.object.organizer, self.object.guest_name)
            self.object.save(update_fields=["welcome_message", "updated_at"])
        messages.success(self.request, "Invitation mise a jour.")
        return reverse("invitations:detail", kwargs={"slug": self.object.slug})


class InvitationDeleteView(InvitationPermissionMixin, View):
    def post(self, request, slug):
        invitation = get_object_or_404(self.get_queryset(), slug=slug)
        if invitation.is_locked:
            messages.error(
                request,
                "Suppression impossible: l'invitation a deja ete imprimee ou partagee.",
            )
            return redirect("invitations:detail", slug=slug)
        invitation.delete()
        messages.success(request, "Invitation supprimee.")
        return redirect("invitations:list")


class SharedInvitationDetailView(DetailView):
    model = Invitation
    template_name = "invitations/shared_invitation_detail.html"
    context_object_name = "invitation"
    slug_field = "share_token"
    slug_url_kwarg = "token"

    def get_context_data(self, **kwargs):
        _ensure_invitation_qr(self.object)
        return super().get_context_data(**kwargs)


class StaffReportView(StaffRequiredMixin, TemplateView):
    template_name = "invitations/staff_report.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        profiles = OrganizerProfile.objects.select_related("user").prefetch_related("subscriptions")
        subscriptions = Subscription.objects.select_related("organizer", "organizer__user")
        transactions = PaymentTransaction.objects.select_related(
            "subscription", "subscription__organizer", "subscription__organizer__user"
        )[:20]
        support_threads = list(
            OrganizerProfile.objects.select_related("user")
            .annotate(
                latest_support_at=Max("support_messages__created_at"),
                unread_for_admin_count=Count(
                    "support_messages",
                    filter=Q(
                        support_messages__sender_type=SupportMessage.SENDER_USER,
                        support_messages__is_read_by_admin=False,
                    ),
                ),
            )
            .filter(latest_support_at__isnull=False)
            .prefetch_related(
                Prefetch(
                    "support_messages",
                    queryset=SupportMessage.objects.select_related("sender_user").order_by("created_at"),
                )
            )
            .order_by("-latest_support_at", "user__username")
        )
        for thread in support_threads:
            thread.thread_messages = list(thread.support_messages.all())
            thread.latest_support_message = thread.thread_messages[-1] if thread.thread_messages else None
        connected_profiles = _connected_profiles_queryset()
        chart_data = build_staff_stats()
        context.update(
            {
                "profiles": profiles,
                "subscriptions": subscriptions,
                "transactions": transactions,
                "support_threads": support_threads,
                "connected_profiles": connected_profiles,
                "chart_data": chart_data,
                "totals": {
                    "users": User.objects.count(),
                    "organizers": profiles.count(),
                    "subscriptions": subscriptions.count(),
                    "paid_subscriptions": subscriptions.filter(is_paid=True).count(),
                    "active_users": chart_data["activity"][0]["total"] if chart_data["activity"] else 0,
                    "inactive_users": chart_data["activity"][1]["total"] if len(chart_data["activity"]) > 1 else 0,
                    "revenue_month": chart_data["revenue"][0]["total"] if chart_data["revenue"] else 0,
                    "site_hits": chart_data["visit_totals"]["total_hits"],
                    "unique_visitors": chart_data["visit_totals"]["unique_visitors"],
                    "connected_users": connected_profiles.count(),
                },
            }
        )
        return context


class StaffReportExcelView(StaffRequiredMixin, View):
    def get(self, request):
        workbook = build_staff_report_workbook()
        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="rapport-utilisateurs-abonnements.xlsx"'
        return response


class ContactAdminView(LoginRequiredMixin, TemplateView):
    template_name = "invitations/contact_admin.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context["form"] = kwargs.get("form") or ContactAdminForm()
        context["reply_form"] = kwargs.get("reply_form") or ContactAdminReplyForm()
        profile = getattr(self.request.user, "organizer_profile", None)
        if profile:
            _mark_admin_messages_read_by_user(profile)
            thread_messages = list(
                profile.support_messages.select_related("sender_user").order_by("created_at")[:50]
            )
            context["messages_thread"] = thread_messages
            context["latest_thread_message"] = thread_messages[-1] if thread_messages else None
            context["thread_started"] = profile.support_messages.exists()
            context["can_send_message"] = profile.support_user_can_send
        return context

    def post(self, request, *args, **kwargs):
        form = ContactAdminForm(request.POST)
        if form.is_valid():
            profile = getattr(request.user, "organizer_profile", None)
            if not profile:
                messages.error(request, "Votre profil organisateur est incomplet.")
                return redirect("invitations:home")
            if not profile.support_user_can_send:
                messages.error(request, "L'envoi de messages a ete desactive par l'administration.")
                return redirect("invitations:contact-admin")
            if profile.support_messages.exists():
                messages.error(request, "Une discussion est deja ouverte. Utilisez le bouton repondre.")
                return redirect("invitations:contact-admin")
            subject = f"[Contact utilisateur] {form.cleaned_data['subject']}"
            body = (
                f"Utilisateur: {request.user.get_full_name() or request.user.username}\n"
                f"Email: {request.user.email or 'non renseigne'}\n"
                f"Ceremonie: {profile.get_event_label() if profile else 'non renseignee'}\n"
                f"Telephone: {profile.phone_number if profile else '-'}\n\n"
                f"{form.cleaned_data['message']}"
            )
            send_mail(
                subject,
                body,
                settings.DEFAULT_FROM_EMAIL,
                [settings.ADMIN_CONTACT_EMAIL],
                fail_silently=True,
            )
            support_message = SupportMessage.objects.create(
                organizer=profile,
                sender_type=SupportMessage.SENDER_USER,
                sender_user=request.user,
                subject=form.cleaned_data["subject"],
                message=form.cleaned_data["message"],
                is_read_by_admin=False,
                is_read_by_user=True,
            )
            _broadcast_support_message(support_message)
            profile.support_thread_subject = form.cleaned_data["subject"]
            profile.save(update_fields=["support_thread_subject"])
            messages.success(request, "Votre message a ete envoye a l'administration.")
            if _is_ajax_request(request):
                return JsonResponse(
                    {
                        "ok": True,
                        "message_item": _serialize_support_message(support_message, for_staff=False),
                    }
                )
            return redirect("invitations:contact-admin")
        if _is_ajax_request(request):
            return JsonResponse({"ok": False, "errors": form.errors}, status=400)
        return self.render_to_response(self.get_context_data(form=form))


class ContactAdminReplyView(LoginRequiredMixin, View):
    def post(self, request):
        profile = getattr(request.user, "organizer_profile", None)
        if not profile:
            messages.error(request, "Votre profil organisateur est incomplet.")
            return redirect("invitations:home")
        if not profile.support_messages.exists():
            messages.error(request, "Aucune discussion n'a encore ete ouverte.")
            return redirect("invitations:contact-admin")
        if not profile.support_user_can_send:
            messages.error(request, "L'administration a desactive vos reponses.")
            return redirect("invitations:contact-admin")

        form = ContactAdminReplyForm(request.POST)
        if not form.is_valid():
            messages.error(request, "Merci de verifier votre reponse avant envoi.")
            return redirect("invitations:contact-admin")

        support_message = SupportMessage.objects.create(
            organizer=profile,
            sender_type=SupportMessage.SENDER_USER,
            sender_user=request.user,
            subject=profile.support_thread_subject or "Suite de la discussion",
            message=form.cleaned_data["message"],
            is_read_by_admin=False,
            is_read_by_user=True,
        )
        _broadcast_support_message(support_message)
        send_mail(
            f"[Suite utilisateur] {profile.support_thread_subject or 'Suite de la discussion'}",
            form.cleaned_data["message"],
            settings.DEFAULT_FROM_EMAIL,
            [settings.ADMIN_CONTACT_EMAIL],
            fail_silently=True,
        )
        messages.success(request, "Votre reponse a ete envoyee a l'administration.")
        if _is_ajax_request(request):
            return JsonResponse(
                {
                    "ok": True,
                    "message_item": _serialize_support_message(support_message, for_staff=False),
                }
            )
        return redirect("invitations:contact-admin")


class StaffReplySupportView(StaffRequiredMixin, View):
    def post(self, request, organizer_id):
        profile = get_object_or_404(OrganizerProfile, pk=organizer_id)
        form = AdminSupportReplyForm(request.POST)
        if not form.is_valid():
            messages.error(request, "La reponse admin est invalide.")
            if _is_ajax_request(request):
                return JsonResponse({"ok": False, "errors": form.errors}, status=400)
            return redirect(f"{reverse('invitations:staff-report')}#messages")

        read_ids = list(
            profile.support_messages.filter(
                sender_type=SupportMessage.SENDER_USER,
                is_read_by_admin=False,
            ).values_list("id", flat=True)
        )
        profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_USER,
            is_read_by_admin=False,
        ).update(is_read_by_admin=True)
        if read_ids:
            _broadcast_support_read_state(profile, read_by_admin_ids=read_ids)
        support_message = SupportMessage.objects.create(
            organizer=profile,
            sender_type=SupportMessage.SENDER_ADMIN,
            sender_user=request.user,
            subject=form.cleaned_data["subject"],
            message=form.cleaned_data["message"],
            is_read_by_admin=True,
            is_read_by_user=False,
        )
        _broadcast_support_message(support_message)
        send_mail(
            f"[Administration] {form.cleaned_data['subject']}",
            form.cleaned_data["message"],
            settings.DEFAULT_FROM_EMAIL,
            [profile.user.email] if profile.user.email else [settings.ADMIN_CONTACT_EMAIL],
            fail_silently=True,
        )
        messages.success(request, "La reponse a ete enregistree et envoyee.")
        if _is_ajax_request(request):
            return JsonResponse(
                {
                    "ok": True,
                    "message_item": _serialize_support_message(support_message, for_staff=True),
                }
            )
        return redirect(f"{reverse('invitations:staff-report')}#messages")


class StaffToggleSupportView(StaffRequiredMixin, View):
    def post(self, request, organizer_id):
        profile = get_object_or_404(OrganizerProfile, pk=organizer_id)
        profile.support_user_can_send = not profile.support_user_can_send
        profile.save(update_fields=["support_user_can_send"])
        _broadcast_support_permissions(profile)
        messages.success(
            request,
            "L'envoi de messages utilisateur a ete active."
            if profile.support_user_can_send
            else "L'envoi de messages utilisateur a ete desactive.",
        )
        if _is_ajax_request(request):
            return JsonResponse({"ok": True, "can_send_message": profile.support_user_can_send})
        return redirect(f"{reverse('invitations:staff-report')}#messages")


@login_required
@require_POST
def mark_support_thread_read(request, organizer_id):
    if not request.user.is_staff:
        return JsonResponse({"ok": False}, status=403)

    profile = get_object_or_404(OrganizerProfile, pk=organizer_id)
    read_ids = list(
        profile.support_messages.filter(
            sender_type=SupportMessage.SENDER_USER,
            is_read_by_admin=False,
        ).values_list("id", flat=True)
    )
    profile.support_messages.filter(
        sender_type=SupportMessage.SENDER_USER,
        is_read_by_admin=False,
    ).update(is_read_by_admin=True)
    total_unread = SupportMessage.objects.filter(
        sender_type=SupportMessage.SENDER_USER,
        is_read_by_admin=False,
    ).count()
    _broadcast_support_read_state(profile, read_by_admin_ids=read_ids)
    return JsonResponse(
        {
            "ok": True,
            "organizer_id": profile.id,
            "thread_unread_count": 0,
            "unread_count": total_unread,
        }
    )


@login_required
def contact_admin_thread_data(request):
    profile = getattr(request.user, "organizer_profile", None)
    if not profile:
        return JsonResponse({"messages": [], "latest_timestamp": ""})

    _mark_admin_messages_read_by_user(profile)
    thread_messages = list(
        profile.support_messages.select_related("sender_user").order_by("created_at")[:50]
    )
    latest = thread_messages[-1].created_at.isoformat() if thread_messages else ""
    return JsonResponse(
        {
            "messages": [
                _serialize_support_message(item, for_staff=False)
                for item in thread_messages
            ],
            "latest_timestamp": latest,
        }
    )


@login_required
def staff_support_feed(request):
    if not request.user.is_staff:
        return JsonResponse({"items": []}, status=403)

    support_messages = SupportMessage.objects.select_related("organizer", "organizer__user").order_by("-created_at")[:12]
    latest = support_messages[0].created_at.isoformat() if support_messages else ""
    return JsonResponse(
        {
            "items": [
                {
                    "organizer_id": item.organizer_id,
                    "username": item.organizer.user.username,
                    "subject": item.subject,
                    "message": item.message,
                    "sender_label": item.sender_label,
                    "created_at": timezone.localtime(item.created_at).strftime("%d/%m/%Y %H:%M"),
                    "can_send": item.organizer.support_user_can_send,
                }
                for item in support_messages
            ],
            "latest_timestamp": latest,
        }
    )


@login_required
@require_POST
def update_welcome_message(request):
    profile = request.user.organizer_profile
    form = WelcomeMessageForm(request.POST, instance=profile)
    if form.is_valid():
        form.save()
        messages.success(
            request,
            tr_text(
                "Le message de bienvenue global a ete enregistre.",
                "The global welcome message has been saved.",
            ),
        )
    else:
        messages.error(
            request,
            tr_text(
                "Merci de verifier le message de bienvenue.",
                "Please review the welcome message.",
            ),
        )
    return redirect("invitations:subscription")


@login_required
@require_POST
def submit_site_rating(request):
    profile = request.user.organizer_profile
    form = SiteRatingForm(request.POST)
    if form.is_valid():
        SiteRating.objects.update_or_create(
            organizer=profile,
            defaults={"stars": form.cleaned_data["stars"]},
        )
        messages.success(
            request,
            tr_text("Merci pour votre note.", "Thank you for your rating."),
        )
    else:
        messages.error(
            request,
            tr_text("La note envoyee est invalide.", "The submitted rating is invalid."),
        )
    return redirect("invitations:home")


@login_required
def notifications_poll(request):
    organizer = getattr(request.user, "organizer_profile", None)
    if request.user.is_staff:
        unread_count = SupportMessage.objects.filter(sender_type=SupportMessage.SENDER_USER, is_read_by_admin=False).count()
        latest = SupportMessage.objects.filter(sender_type=SupportMessage.SENDER_USER).order_by("-created_at").first()
    elif organizer:
        unread_count = organizer.support_messages.filter(sender_type=SupportMessage.SENDER_ADMIN, is_read_by_user=False).count()
        latest = organizer.support_messages.filter(sender_type=SupportMessage.SENDER_ADMIN).order_by("-created_at").first()
    else:
        unread_count = 0
        latest = None

    return JsonResponse(
        {
            "unread_count": unread_count,
            "latest_timestamp": latest.created_at.isoformat() if latest else "",
        }
    )


@login_required
@require_POST
@transaction.atomic
def start_payment(request):
    profile = OrganizerProfile.objects.select_for_update().get(user=request.user)
    subscription = profile.active_subscription or profile.subscriptions.first()
    active_subscription = profile.active_subscription
    form = PaymentInitiationForm(request.POST)
    if not subscription:
        messages.error(request, "Aucun abonnement a regler.")
        return redirect("invitations:subscription")

    if not form.is_valid():
        messages.error(request, "Merci de verifier les informations de paiement.")
        return redirect("invitations:subscription")

    desired_invitations = form.cleaned_data.get("planned_invitations") or profile.planned_invitations
    target_subscription = subscription
    if active_subscription and active_subscription.is_paid:
        target_subscription = _build_subscription(
            profile,
            form.cleaned_data["provider"],
            invite_count=desired_invitations,
            previous_subscription=active_subscription,
        )
        profile.planned_invitations = desired_invitations
        profile.save(update_fields=["planned_invitations"])
    else:
        plan_code, invitation_limit, price_usd = Subscription.pick_plan(desired_invitations)
        target_subscription.provider = form.cleaned_data["provider"]
        target_subscription.plan_code = plan_code
        target_subscription.invitation_limit = invitation_limit
        target_subscription.price_usd = price_usd
        target_subscription.save(update_fields=["provider", "plan_code", "invitation_limit", "price_usd"])

    profile.phone_number = form.cleaned_data["phone_number"]
    profile.save(update_fields=["phone_number"])

    transaction = PaymentTransaction.objects.create(
        subscription=target_subscription,
        provider=target_subscription.provider,
        amount_usd=target_subscription.price_usd,
        phone_number=profile.phone_number,
        status=PaymentTransaction.STATUS_PENDING,
    )
    payload = _mobile_money_gateway_payload(transaction)
    config = _provider_configuration(target_subscription.provider)
    transaction.request_payload = json.dumps(payload, ensure_ascii=False)

    if target_subscription.provider == Subscription.PROVIDER_DEMO_CARD:
        card_number = form.cleaned_data["card_number"].replace(" ", "")
        transaction.external_reference = f"DEMO-{card_number[-4:]}"
        transaction.status = PaymentTransaction.STATUS_SUCCEEDED
        transaction.response_payload = json.dumps(
            {
                "mode": "demo-card",
                "brand": "Visa test",
                "note": "Paiement local de demonstration valide sans appel externe.",
            },
            ensure_ascii=False,
        )
        transaction.save(
            update_fields=["request_payload", "response_payload", "status", "external_reference", "updated_at"]
        )
        _activate_subscription(target_subscription)
        if active_subscription and active_subscription.is_paid:
            messages.success(
                request,
                "Le nouveau forfait a ete active et le reliquat de l'ancien abonnement a ete cumule.",
            )
        else:
            messages.success(
                request,
                "Paiement demo par carte valide. Votre abonnement est maintenant actif en local.",
            )
        return redirect("invitations:subscription")

    if config["endpoint"] and config["api_key"] and config["api_secret"]:
        payload["signature"] = PaymentTransaction.sign_payload(payload, config["api_secret"])
        transaction.response_payload = json.dumps(
            {
                "mode": "live-ready",
                "note": "Configuration detectee. Branchez ici l'appel HTTPS signe vers le fournisseur Mobile Money.",
            },
            ensure_ascii=False,
        )
        if active_subscription and active_subscription.is_paid:
            messages.info(
                request,
                "Paiement initialise. Si le paiement reussit, le reliquat de l'abonnement actuel sera ajoute au nouveau.",
            )
        else:
            messages.info(
                request,
                "Paiement initialise. Connectez l'API du fournisseur pour recevoir la confirmation automatique.",
            )
    else:
        transaction.response_payload = json.dumps(
            {
                "mode": "sandbox-local",
                "note": "Aucune cle API configuree. Validation locale necessaire via l'admin ou le callback.",
            },
            ensure_ascii=False,
        )
        messages.warning(
            request,
            "Les identifiants API Mobile Money ne sont pas encore configures. "
            "Le paiement reste en attente jusqu'a confirmation.",
        )

    transaction.save(update_fields=["request_payload", "response_payload", "status", "updated_at"])
    return redirect("invitations:subscription")


@csrf_exempt
def payment_callback(request, provider):
    if request.method != "POST":
        return HttpResponse(status=405)

    try:
        payload = json.loads(request.body.decode("utf-8"))
    except json.JSONDecodeError:
        return HttpResponse(status=400)

    reference = payload.get("merchant_reference")
    if not reference:
        return HttpResponse(status=400)

    transaction = get_object_or_404(PaymentTransaction, transaction_reference=reference, provider=provider)
    config = _provider_configuration(provider)
    signature = request.headers.get("X-Signature", "")

    if config["api_secret"]:
        expected_signature = PaymentTransaction.sign_payload(payload, config["api_secret"])
        if not hmac_compare(signature, expected_signature):
            return HttpResponseForbidden("Invalid signature")

    transaction.response_payload = request.body.decode("utf-8")
    transaction.external_reference = payload.get("provider_reference", "")
    transaction.status = (
        PaymentTransaction.STATUS_SUCCEEDED
        if payload.get("status") == "success"
        else PaymentTransaction.STATUS_FAILED
    )
    transaction.save(update_fields=["response_payload", "external_reference", "status", "updated_at"])

    subscription = transaction.subscription
    if transaction.status == PaymentTransaction.STATUS_SUCCEEDED:
        _activate_subscription(subscription)

    return HttpResponse(status=204)


@login_required
@require_POST
@transaction.atomic
def mark_printed(request, slug):
    invitation = get_object_or_404(
        Invitation.objects.select_for_update(), slug=slug, organizer=request.user.organizer_profile
    )
    invitation.mark_printed()
    return redirect(f"{reverse('invitations:detail', kwargs={'slug': slug})}?autoprint=1")


@login_required
@require_POST
@transaction.atomic
def mark_shared(request, slug):
    invitation = get_object_or_404(
        Invitation.objects.select_for_update(), slug=slug, organizer=request.user.organizer_profile
    )
    invitation.mark_shared()
    absolute_url = request.build_absolute_uri(
        reverse("invitations:shared-detail", kwargs={"token": invitation.share_token})
    )
    text = (
        f"Invitation {invitation.event_label} pour {invitation.guest_name}. "
        f"Place: {invitation.seat_location}. Voir l'invitation: {absolute_url}"
    )
    return redirect(f"https://wa.me/?text={quote(text)}")


@login_required
def download_qrcode(request, slug):
    invitation = get_object_or_404(
        Invitation.objects.select_related("subscription"), slug=slug, organizer=request.user.organizer_profile
    )
    _ensure_invitation_qr(invitation)
    if not invitation.subscription.is_paid or not invitation.qr_code:
        raise Http404("QR code indisponible")
    return FileResponse(
        invitation.qr_code.open("rb"),
        as_attachment=True,
        filename=f"qrcode-{invitation.slug}.png",
    )


@login_required
def invitation_qr_preview(request, slug):
    invitation = get_object_or_404(
        Invitation.objects.select_related("subscription"), slug=slug, organizer=request.user.organizer_profile
    )
    _ensure_invitation_qr(invitation)
    if not invitation.qr_code:
        raise Http404("QR code indisponible")
    return FileResponse(invitation.qr_code.open("rb"), content_type="image/png")


def shared_invitation_qr_preview(request, token):
    invitation = get_object_or_404(Invitation.objects.select_related("subscription"), share_token=token)
    _ensure_invitation_qr(invitation)
    if not invitation.qr_code:
        raise Http404("QR code indisponible")
    return FileResponse(invitation.qr_code.open("rb"), content_type="image/png")


@login_required
def download_invitation_image(request, slug):
    invitation = get_object_or_404(
        Invitation.objects.select_related("subscription", "organizer", "organizer__user"),
        slug=slug,
        organizer=request.user.organizer_profile,
    )
    _ensure_invitation_qr(invitation)
    if not invitation.subscription.is_paid:
        raise Http404("Image indisponible")
    jpeg = build_invitation_jpeg(invitation)
    response = HttpResponse(jpeg.getvalue(), content_type="image/jpeg")
    response["Content-Disposition"] = f'attachment; filename="invitation-{invitation.slug}.jpg"'
    return response


@login_required
def download_excel_template(request):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Invites"
    sheet.append(["Nom du couple ou de l'invite", "Emplacement dans la salle"])
    sheet.append(["Couple Kasongo", "Table 4"])
    sheet.append(["Famille Mbayo", "Rang B - Siege 12"])
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    response = HttpResponse(
        output.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    response["Content-Disposition"] = 'attachment; filename="modele-import-invitations.xlsx"'
    return response
